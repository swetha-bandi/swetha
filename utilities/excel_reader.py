import os
from openpyxl import load_workbook

from utilities.logger import Logger


class ExcelReader:

    logger = Logger.get_logger()

    @staticmethod
    def get_data(file_name, sheet_name):

        try:
            file_path = os.path.join(
                os.path.dirname(os.path.dirname(__file__)),
                "testdata",
                file_name
            )

            workbook = load_workbook(file_path)
            sheet = workbook[sheet_name]

            data = []

            headers = [
                cell.value
                for cell in sheet[1]
            ]

            for row in sheet.iter_rows(min_row=2, values_only=True):
                row_data = dict(zip(headers, row))
                data.append(row_data)

            ExcelReader.logger.info(
                f"Successfully loaded {len(data)} records from '{sheet_name}' sheet."
            )

            return data

        except Exception as e:
            ExcelReader.logger.error(
                f"Failed to read Excel file '{file_name}' and sheet '{sheet_name}'. Error: {e}"
            )
            raise









